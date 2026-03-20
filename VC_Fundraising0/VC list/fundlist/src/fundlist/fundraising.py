from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sqlite3
import urllib.error
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from .store import ensure_parent_dir


DEFAULT_FUNDRAISE_FILES = [
    "/Users/ej/Downloads/2025-2026 Fund Raising - Web2 VC.csv",
    "/Users/ej/Downloads/2025-2026 Fund Raising - Web3 VC.csv",
    "/Users/ej/Downloads/2025-2026 Fund Raising - grants program.csv",
    "/Users/ej/Downloads/2025-2026 Fund Raising - Accelerator Program.csv",
    "/Users/ej/Downloads/2025-2026 Fund Raising - contact.csv",
    "/Users/ej/Downloads/2025-2026 Fund Raising - Email From Pitchdeck.tsv",
    "/Users/ej/Downloads/2025-2026 Fund Raising.xlsx",
]


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
URL_RE = re.compile(r"https?://[^\s,)\]]+")


@dataclass
class FundraisingRecord:
    source_file: str
    source_row: int
    category: str
    org_name: str
    contact_name: str
    email: str
    website: str
    status: str
    region: str
    funding: str
    date_text: str
    notes: str
    raw_json: str

    @property
    def fingerprint(self) -> str:
        raw = "|".join(
            [
                self.source_file,
                str(self.source_row),
                self.category,
                self.org_name,
                self.contact_name,
                self.email,
                self.website,
                self.status,
                self.region,
                self.funding,
                self.date_text,
                self.notes,
            ]
        )
        return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def now_utc_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def normalize_text(value: str) -> str:
    return re.sub(r"[^0-9a-zA-Z가-힣]+", "", value.strip().lower())


def sanitize(value: str, limit: int = 1000) -> str:
    return value.strip().replace("\x00", "")[:limit]


def category_from_path(path: Path) -> str:
    name = path.name.lower()
    if "web2 vc" in name:
        return "web2_vc"
    if "web3 vc" in name:
        return "web3_vc"
    if "grants program" in name:
        return "grants_program"
    if "accelerator program" in name:
        return "accelerator_program"
    if "contact" in name:
        return "vc_contact"
    if "email from pitchdeck" in name:
        return "pitchdeck_email"
    if name.endswith(".xlsx"):
        return "xlsx_sheet"
    return "unknown"


class FundraisingStore:
    def __init__(self, db_path: str) -> None:
        ensure_parent_dir(db_path)
        self.conn = sqlite3.connect(db_path)
        self._init_schema()

    def _init_schema(self) -> None:
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS fundraising_records (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              source_file TEXT NOT NULL,
              source_row INTEGER NOT NULL,
              category TEXT NOT NULL,
              org_name TEXT NOT NULL,
              contact_name TEXT NOT NULL,
              email TEXT NOT NULL,
              website TEXT NOT NULL,
              status TEXT NOT NULL,
              region TEXT NOT NULL,
              funding TEXT NOT NULL,
              date_text TEXT NOT NULL,
              notes TEXT NOT NULL,
              raw_json TEXT NOT NULL,
              fingerprint TEXT NOT NULL UNIQUE,
              imported_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
            """
        )
        self.conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_fundraise_cat ON fundraising_records(category, imported_at DESC)"
        )
        self.conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_fundraise_email ON fundraising_records(email)"
        )
        self.conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_fundraise_org ON fundraising_records(org_name)"
        )
        self.conn.commit()

    def insert_records(self, records: Sequence[FundraisingRecord]) -> int:
        sql = """
            INSERT OR IGNORE INTO fundraising_records (
              source_file, source_row, category, org_name, contact_name, email, website,
              status, region, funding, date_text, notes, raw_json, fingerprint
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """
        before = self.conn.total_changes
        with self.conn:
            self.conn.executemany(
                sql,
                [
                    (
                        r.source_file,
                        r.source_row,
                        r.category,
                        r.org_name,
                        r.contact_name,
                        r.email,
                        r.website,
                        r.status,
                        r.region,
                        r.funding,
                        r.date_text,
                        r.notes,
                        r.raw_json,
                        r.fingerprint,
                    )
                    for r in records
                ],
            )
        return self.conn.total_changes - before

    def stats(self) -> Dict[str, object]:
        cur = self.conn.cursor()
        total = cur.execute("SELECT COUNT(*) FROM fundraising_records").fetchone()[0]
        uniq_orgs = cur.execute(
            "SELECT COUNT(DISTINCT org_name) FROM fundraising_records WHERE org_name <> ''"
        ).fetchone()[0]
        uniq_emails = cur.execute(
            "SELECT COUNT(DISTINCT email) FROM fundraising_records WHERE email <> ''"
        ).fetchone()[0]
        by_category = cur.execute(
            """
            SELECT category, COUNT(*) AS n
            FROM fundraising_records
            GROUP BY category
            ORDER BY n DESC
            """
        ).fetchall()
        missing_email = cur.execute(
            """
            SELECT category, COUNT(*) AS n
            FROM fundraising_records
            WHERE email = ''
            GROUP BY category
            ORDER BY n DESC
            """
        ).fetchall()
        top_targets = cur.execute(
            """
            SELECT category, org_name, email, website, notes
            FROM fundraising_records
            WHERE (email <> '' OR website <> '') AND org_name <> ''
            ORDER BY category, org_name
            LIMIT 80
            """
        ).fetchall()
        return {
            "total_records": total,
            "unique_orgs": uniq_orgs,
            "unique_emails": uniq_emails,
            "by_category": [{"category": row[0], "count": row[1]} for row in by_category],
            "missing_email_by_category": [{"category": row[0], "count": row[1]} for row in missing_email],
            "top_targets": [
                {
                    "category": row[0],
                    "org_name": row[1],
                    "email": row[2],
                    "website": row[3],
                    "notes": (row[4] or "")[:200],
                }
                for row in top_targets
            ],
        }


def extract_emails(text: str) -> List[str]:
    return sorted(set(email.lower() for email in EMAIL_RE.findall(text or "")))


def extract_urls(text: str) -> List[str]:
    return sorted(set(URL_RE.findall(text or "")))


def detect_header_index(rows: Sequence[List[str]]) -> int:
    hints = [
        "fund",
        "task",
        "program",
        "website",
        "link",
        "date",
        "amount",
        "tier",
        "region",
        "active",
        "notes",
        "현재상태",
        "투자사",
        "이메일",
        "그랜트",
        "조직",
        "펀딩규모",
        "금액",
        "설명",
    ]
    best_idx = 0
    best_score = -1
    for idx, row in enumerate(rows):
        cleaned = [c.strip() for c in row]
        non_empty = [c for c in cleaned if c]
        if len(non_empty) < 2:
            continue
        row_join = " ".join(cleaned).lower()
        score = sum(1 for h in hints if h in row_join) + len(non_empty) // 4
        if score > best_score:
            best_score = score
            best_idx = idx
    return best_idx


def map_row(
    category: str,
    source_file: str,
    source_row: int,
    headers: Sequence[str],
    row: Sequence[str],
) -> Optional[FundraisingRecord]:
    mapping = {normalize_text(h): sanitize(row[i]) if i < len(row) else "" for i, h in enumerate(headers)}
    row_join = " | ".join([sanitize(v, limit=400) for v in row if v.strip()])

    def pick(*aliases: str) -> str:
        for alias in aliases:
            a = normalize_text(alias)
            for k, v in mapping.items():
                if not v:
                    continue
                if k == a or a in k:
                    return v
        return ""

    org_name = pick("fundname", "taskname", "program", "투자사", "조직", "그랜트이름", "name")
    if org_name.isdigit():
        org_name = ""
    contact_name = pick("contactname", "담당자", "이름")
    email_field = pick("email", "이메일", "이메일연락처", "contact", "연락처")
    website = pick("website", "link", "url", "링크")
    status = pick("현재상태", "status", "active", "상태날짜")
    region = pick("region", "location", "분야", "지역")
    funding = pick("amountraisedmusd", "amountraised", "펀딩규모", "금액", "amount")
    date_text = pick("dateannouncedandraised", "date", "상태날짜")
    notes = pick("notes", "설명", "출처", "결과", "제출내용적은링크")

    emails = extract_emails(email_field or row_join)
    urls = extract_urls(website or email_field or row_join)
    email = ",".join(emails[:3]) if emails else ""
    if not website and urls:
        website = urls[0]

    if not any([org_name, email, website, notes, funding, date_text]):
        return None

    return FundraisingRecord(
        source_file=source_file,
        source_row=source_row,
        category=category,
        org_name=sanitize(org_name, limit=240),
        contact_name=sanitize(contact_name, limit=240),
        email=sanitize(email, limit=240),
        website=sanitize(website, limit=500),
        status=sanitize(status, limit=200),
        region=sanitize(region, limit=120),
        funding=sanitize(funding, limit=200),
        date_text=sanitize(date_text, limit=120),
        notes=sanitize(notes or row_join, limit=800),
        raw_json=json.dumps({"headers": list(headers), "row": list(row)}, ensure_ascii=False),
    )


def parse_delimited_file(path: Path) -> List[List[str]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows: List[List[str]] = []
    with path.open("r", newline="", encoding="utf-8-sig") as fp:
        reader = csv.reader(fp, delimiter=delimiter)
        for row in reader:
            rows.append([cell.strip() for cell in row])
    return rows


def parse_pitchdeck_tsv(path: Path) -> List[FundraisingRecord]:
    rows = parse_delimited_file(path)
    out: List[FundraisingRecord] = []
    seen: set[str] = set()
    for idx, row in enumerate(rows, start=1):
        text = " ".join(row)
        for email in extract_emails(text):
            if email in seen:
                continue
            seen.add(email)
            out.append(
                FundraisingRecord(
                    source_file=str(path),
                    source_row=idx,
                    category="pitchdeck_email",
                    org_name="",
                    contact_name="",
                    email=email,
                    website="",
                    status="",
                    region="",
                    funding="",
                    date_text="",
                    notes="email list imported from pitchdeck tsv",
                    raw_json=json.dumps({"row": row}, ensure_ascii=False),
                )
            )
    return out


def parse_csv_like(path: Path) -> List[FundraisingRecord]:
    category = category_from_path(path)
    if category == "pitchdeck_email":
        return parse_pitchdeck_tsv(path)

    rows = parse_delimited_file(path)
    if not rows:
        return []

    header_idx = detect_header_index(rows)
    headers = rows[header_idx] if header_idx < len(rows) else []
    headers = [h if h else f"col_{i+1}" for i, h in enumerate(headers)]

    out: List[FundraisingRecord] = []
    for i, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not any(cell.strip() for cell in row):
            continue
        rec = map_row(
            category=category,
            source_file=str(path),
            source_row=i,
            headers=headers,
            row=row,
        )
        if rec:
            out.append(rec)
    return out


def parse_xlsx(path: Path) -> List[FundraisingRecord]:
    try:
        import openpyxl  # type: ignore
    except Exception:  # noqa: BLE001
        print(f"[warn] openpyxl is not installed; skipping xlsx file: {path}")
        return []

    out: List[FundraisingRecord] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows: List[List[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(cell).strip() if cell is not None else "" for cell in row])
        if not rows:
            continue
        header_idx = detect_header_index(rows)
        headers = [h if h else f"col_{i+1}" for i, h in enumerate(rows[header_idx])]
        for i, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            if not any(cell.strip() for cell in row):
                continue
            rec = map_row(
                category=f"xlsx:{sheet}",
                source_file=str(path),
                source_row=i,
                headers=headers,
                row=row,
            )
            if rec:
                out.append(rec)
    return out


def import_fundraising_files(db_path: str, files: Sequence[str]) -> Tuple[int, int]:
    store = FundraisingStore(db_path)
    all_records: List[FundraisingRecord] = []

    for raw in files:
        p = Path(raw).expanduser()
        if not p.exists():
            print(f"[warn] file not found: {p}")
            continue
        try:
            if p.suffix.lower() in {".csv", ".tsv"}:
                parsed = parse_csv_like(p)
            elif p.suffix.lower() == ".xlsx":
                parsed = parse_xlsx(p)
            else:
                print(f"[warn] unsupported extension skipped: {p}")
                continue
        except Exception as exc:  # noqa: BLE001
            print(f"[warn] parse failed: {p} -> {exc}")
            continue

        all_records.extend(parsed)
        print(f"[ok] imported from {p.name}: {len(parsed)}")

    inserted = store.insert_records(all_records)
    return len(all_records), inserted


def get_hf_token() -> str:
    return (
        os.environ.get("HF_TOKEN", "").strip()
        or os.environ.get("HUGGINGFACE_API_KEY", "").strip()
        or os.environ.get("HUGGINGFACEHUB_API_TOKEN", "").strip()
    )


def get_openrouter_key() -> str:
    return os.environ.get("OPENROUTER_API_KEY", "").strip()


def call_groq_summary(stats: Dict[str, object], model: str) -> str:
    api_key = os.environ.get("GROQ_API_KEY", "").strip()
    if not api_key:
        return "GROQ_API_KEY is not set. AI summary skipped."

    endpoint = os.environ.get("GROQ_BASE_URL", "https://api.groq.com/openai/v1/chat/completions")
    prompt = (
        "You are a fundraising operations analyst. "
        "Write a concise Korean report with: (1) 핵심 인사이트 5개, "
        "(2) 2주 실행 계획, (3) 우선 연락 대상 10개. "
        "Use the provided JSON stats only.\n\n"
        + json.dumps(stats, ensure_ascii=False)
    )
    payload = {
        "model": model,
        "temperature": 0.2,
        "messages": [
            {"role": "system", "content": "Return practical, execution-ready Korean guidance."},
            {"role": "user", "content": prompt},
        ],
    }

    req = urllib.request.Request(
        endpoint,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            body = resp.read().decode("utf-8")
    except Exception as exc:  # noqa: BLE001
        return f"AI summary failed: {exc}"

    try:
        parsed = json.loads(body)
        return (
            parsed.get("choices", [{}])[0]
            .get("message", {})
            .get("content", "AI returned empty response.")
        )
    except Exception as exc:  # noqa: BLE001
        return f"AI summary parse failed: {exc}"


def call_gemini_summary(stats: Dict[str, object], model: str) -> str:
    api_key = os.environ.get("GEMINI_API_KEY", "").strip()
    if not api_key:
        return "GEMINI_API_KEY is not set. AI summary skipped."

    endpoint = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
    prompt = (
        "You are a fundraising operations analyst. "
        "Write a concise Korean report with: (1) 핵심 인사이트 5개, "
        "(2) 2주 실행 계획, (3) 우선 연락 대상 10개. "
        "Use the provided JSON stats only.\n\n"
        + json.dumps(stats, ensure_ascii=False)
    )
    payload = {
        "contents": [
            {
                "role": "user",
                "parts": [{"text": prompt}],
            }
        ],
        "generationConfig": {"temperature": 0.2},
    }

    req = urllib.request.Request(
        endpoint,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            body = resp.read().decode("utf-8")
    except Exception as exc:  # noqa: BLE001
        return f"AI summary failed: {exc}"

    try:
        parsed = json.loads(body)
        candidates = parsed.get("candidates", [])
        if not candidates:
            return f"AI summary failed: {parsed}"
        content = candidates[0].get("content", {})
        parts = content.get("parts", []) if isinstance(content, dict) else []
        if parts and isinstance(parts[0], dict) and parts[0].get("text"):
            return str(parts[0]["text"])
        return "AI returned empty response."
    except Exception as exc:  # noqa: BLE001
        return f"AI summary parse failed: {exc}"


def call_huggingface_summary(stats: Dict[str, object], model: str) -> str:
    api_key = get_hf_token()
    if not api_key:
        return "HF_TOKEN (or HUGGINGFACE_API_KEY) is not set. AI summary skipped."

    endpoint = os.environ.get("HUGGINGFACE_BASE_URL", "https://router.huggingface.co/v1/chat/completions")
    prompt = (
        "You are a fundraising operations analyst. "
        "Write a concise Korean report with: (1) 핵심 인사이트 5개, "
        "(2) 2주 실행 계획, (3) 우선 연락 대상 10개. "
        "Use the provided JSON stats only.\n\n"
        + json.dumps(stats, ensure_ascii=False)
    )
    router_payload = {
        "model": model,
        "temperature": 0.2,
        "messages": [
            {"role": "system", "content": "Return practical, execution-ready Korean guidance."},
            {"role": "user", "content": prompt},
        ],
    }

    try:
        req = urllib.request.Request(
            endpoint,
            data=json.dumps(router_payload).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=60) as resp:
            body = resp.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        detail = ""
        try:
            raw = exc.read().decode("utf-8", errors="ignore")
            detail = re.sub(r"<[^>]+>", " ", raw)
            detail = " ".join(detail.split())[:220]
        except Exception:  # noqa: BLE001
            pass
        return f"AI summary failed: HTTP {exc.code}{' - ' + detail if detail else ''}"
    except Exception as exc:  # noqa: BLE001
        return f"AI summary failed: {exc}"

    try:
        parsed = json.loads(body)
        if isinstance(parsed, dict):
            choices = parsed.get("choices")
            if isinstance(choices, list) and choices:
                text = (
                    choices[0].get("message", {}).get("content")
                    if isinstance(choices[0], dict)
                    else None
                )
                if text:
                    return str(text)
            if parsed.get("generated_text"):
                return str(parsed["generated_text"])
            if parsed.get("error"):
                return f"AI summary failed: {parsed.get('error')}"
        if isinstance(parsed, list) and parsed and isinstance(parsed[0], dict):
            if parsed[0].get("generated_text"):
                return str(parsed[0]["generated_text"])
        return "AI returned empty response."
    except Exception as exc:  # noqa: BLE001
        return f"AI summary parse failed: {exc}"


def call_openrouter_summary(stats: Dict[str, object], model: str) -> str:
    api_key = get_openrouter_key()
    if not api_key:
        return "OPENROUTER_API_KEY is not set. AI summary skipped."

    endpoint = os.environ.get("OPENROUTER_BASE_URL", "https://openrouter.ai/api/v1/chat/completions")
    prompt = (
        "You are a fundraising operations analyst. "
        "Write a concise Korean report with: (1) 핵심 인사이트 5개, "
        "(2) 2주 실행 계획, (3) 우선 연락 대상 10개. "
        "Use the provided JSON stats only.\n\n"
        + json.dumps(stats, ensure_ascii=False)
    )
    payload = {
        "model": model,
        "temperature": 0.2,
        "messages": [
            {"role": "system", "content": "Return practical, execution-ready Korean guidance."},
            {"role": "user", "content": prompt},
        ],
    }
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "HTTP-Referer": os.environ.get("OPENROUTER_HTTP_REFERER", "https://local.fundlist"),
        "X-Title": os.environ.get("OPENROUTER_APP_TITLE", "fundlist"),
    }

    try:
        req = urllib.request.Request(
            endpoint,
            data=json.dumps(payload).encode("utf-8"),
            headers=headers,
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=60) as resp:
            body = resp.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        detail = ""
        try:
            raw = exc.read().decode("utf-8", errors="ignore")
            detail = re.sub(r"<[^>]+>", " ", raw)
            detail = " ".join(detail.split())[:220]
        except Exception:  # noqa: BLE001
            pass
        return f"AI summary failed: HTTP {exc.code}{' - ' + detail if detail else ''}"
    except Exception as exc:  # noqa: BLE001
        return f"AI summary failed: {exc}"

    try:
        parsed = json.loads(body)
        if isinstance(parsed, dict):
            choices = parsed.get("choices")
            if isinstance(choices, list) and choices and isinstance(choices[0], dict):
                text = choices[0].get("message", {}).get("content")
                if text:
                    return str(text)
            if parsed.get("error"):
                return f"AI summary failed: {parsed.get('error')}"
        return "AI returned empty response."
    except Exception as exc:  # noqa: BLE001
        return f"AI summary parse failed: {exc}"


def write_markdown_report(
    db_path: str,
    output_path: str,
    with_ai: bool,
    model: str,
    ai_provider: str,
) -> str:
    store = FundraisingStore(db_path)
    stats = store.stats()
    out = Path(output_path).expanduser()
    ensure_parent_dir(str(out))

    lines = [
        "# Fundraising Intelligence Report",
        "",
        f"- Generated (UTC): {now_utc_iso()}",
        f"- DB: {db_path}",
        "",
        "## Summary",
        f"- Total records: {stats['total_records']}",
        f"- Unique orgs: {stats['unique_orgs']}",
        f"- Unique emails: {stats['unique_emails']}",
        "",
        "## Category Breakdown",
    ]
    for row in stats["by_category"]:
        lines.append(f"- {row['category']}: {row['count']}")

    lines.extend(["", "## Missing Email By Category"])
    missing = stats["missing_email_by_category"]
    if missing:
        for row in missing:
            lines.append(f"- {row['category']}: {row['count']}")
    else:
        lines.append("- none")

    lines.extend(["", "## Top Outreach Targets"])
    targets = stats["top_targets"][:30]
    if targets:
        for t in targets:
            lines.append(
                f"- [{t['category']}] {t['org_name']} | email={t['email'] or '-'} | site={t['website'] or '-'}"
            )
    else:
        lines.append("- no targets")

    if with_ai:
        provider = ai_provider.strip().lower()
        if provider == "gemini":
            ai_text = call_gemini_summary(stats, model=model)
        elif provider == "huggingface":
            ai_text = call_huggingface_summary(stats, model=model)
        elif provider == "openrouter":
            ai_text = call_openrouter_summary(stats, model=model)
        else:
            ai_text = call_groq_summary(stats, model=model)
            provider = "groq"
        lines.extend(["", f"## AI Summary ({provider})", ai_text])

    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return str(out)


def parse_files_argument(files_arg: str) -> List[str]:
    if not files_arg.strip():
        return []
    return [item.strip() for item in files_arg.split(",") if item.strip()]


def resolve_ai_model(provider: str, model: str) -> str:
    raw = (model or "").strip()
    if raw:
        return raw
    p = provider.strip().lower()
    if p == "gemini":
        return os.environ.get("GEMINI_MODEL", "gemini-2.0-flash")
    if p == "huggingface":
        return os.environ.get("HUGGINGFACE_MODEL", "Qwen/Qwen2.5-7B-Instruct")
    if p == "openrouter":
        return os.environ.get("OPENROUTER_MODEL", "openai/gpt-4o-mini")
    return os.environ.get("GROQ_MODEL", "llama-3.3-70b-versatile")


def fundraise_import_command(args: argparse.Namespace) -> int:
    files = parse_files_argument(args.files)
    fetched, inserted = import_fundraising_files(args.db, files)
    print(f"[done] fundraising import parsed={fetched} inserted={inserted} db={args.db}")
    return 0


def fundraise_report_command(args: argparse.Namespace) -> int:
    model = resolve_ai_model(args.ai_provider, args.model)
    path = write_markdown_report(
        db_path=args.db,
        output_path=args.output,
        with_ai=args.with_ai,
        model=model,
        ai_provider=args.ai_provider,
    )
    print(f"[done] report written: {path}")
    return 0


def fundraise_run_command(args: argparse.Namespace) -> int:
    files = parse_files_argument(args.files)
    fetched, inserted = import_fundraising_files(args.db, files)
    print(f"[ok] fundraising import parsed={fetched} inserted={inserted}")
    model = resolve_ai_model(args.ai_provider, args.model)
    path = write_markdown_report(
        db_path=args.db,
        output_path=args.output,
        with_ai=args.with_ai,
        model=model,
        ai_provider=args.ai_provider,
    )
    print(f"[done] report written: {path}")
    return 0
